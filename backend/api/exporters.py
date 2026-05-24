"""
Export utilities — DOCX and PDF generation from PRDSpec.
Single clean() function used by both exporters.
"""
from __future__ import annotations
import io
from ..models.schemas import PRDSpec


def clean(text: str) -> str:
    """Replace unicode chars that crash legacy fonts."""
    return (str(text)
        .replace('\u2014', '-').replace('\u2013', '-')
        .replace('\u2018', "'").replace('\u2019', "'")
        .replace('\u201c', '"').replace('\u201d', '"')
        .replace('\u2022', '-').replace('\u2026', '...')
        .replace('\u2192', '->').replace('\u00b7', '-')
        .encode('ascii', 'replace').decode('ascii')
    )


def prd_to_docx(prd: PRDSpec) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    def h1(text):
        p = doc.add_heading(clean(text), level=1)
        for run in p.runs:
            run.font.color.rgb = RGBColor(27, 46, 75)

    def h2(text):
        p = doc.add_heading(clean(text), level=2)
        for run in p.runs:
            run.font.color.rgb = RGBColor(51, 65, 85)

    def body(text):
        doc.add_paragraph(clean(text))

    def bullet(text):
        doc.add_paragraph(clean(text), style='List Bullet')

    title_p = doc.add_heading(clean(prd.title), level=0)
    for run in title_p.runs:
        run.font.color.rgb = RGBColor(27, 46, 75)
    doc.add_paragraph(clean(prd.tldr))
    doc.add_paragraph()

    h1("Problem Statement")
    body(prd.problem_statement)
    h1("Goal")
    body(prd.goal)
    h1("Target Users")
    for u in prd.target_users:
        bullet(u)

    if prd.competitive_context:
        h1("Competitive Context")
        for comp in prd.competitive_context:
            h2(comp.competitor)
            body("How they solve it: " + comp.how_they_solve_it)
            body("Gap / opportunity: " + comp.gap_or_opportunity)

    h1("User Stories")
    priority_map = {"must-have": "Must-have", "should-have": "Should-have", "nice-to-have": "Nice-to-have"}
    for i, story in enumerate(prd.user_stories, 1):
        h2(f"Story {i} - {priority_map.get(story.priority, story.priority)}")
        body(f"As a {story.persona}, I want {story.goal} so that {story.benefit}")
        doc.add_paragraph("Acceptance Criteria:")
        for ac in story.criteria:
            bullet(f"Given {ac.given} - when {ac.when} - then {ac.then}")
        doc.add_paragraph()

    h1("Success Metrics")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, txt in zip(table.rows[0].cells, ["Metric", "Baseline", "Target", "Timeframe"]):
        cell.text = txt
        cell.paragraphs[0].runs[0].bold = True
    for m in prd.success_metrics:
        row = table.add_row().cells
        for cell, txt in zip(row, [m.metric, m.baseline, m.target, m.timeframe]):
            cell.text = clean(str(txt))
    doc.add_paragraph()

    h1("Out of Scope")
    for o in prd.out_of_scope:
        bullet(o)
    h1("Edge Cases")
    for e in prd.edge_cases:
        bullet(e)
    h1("Risks")
    for r in prd.risks:
        bullet(r)
    if prd.open_questions:
        h1("Open Questions")
        for q in prd.open_questions:
            bullet(f"[ ] {q}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def prd_to_pdf(prd: PRDSpec) -> bytes:
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 9)
            self.set_text_color(100, 116, 139)
            self.cell(0, 8, clean(prd.title)[:80], align="L")
            self.ln(2)
            self.set_draw_color(203, 213, 225)
            self.set_line_width(0.3)
            self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
            self.ln(4)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(148, 163, 184)
            self.cell(0, 10, f"Page {self.page_no()}", align="C")

    pdf = PDF()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    def h1(text):
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 16)
        pdf.set_text_color(27, 46, 75)
        pdf.multi_cell(0, 8, clean(text))
        pdf.set_draw_color(37, 99, 235)
        pdf.set_line_width(0.5)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
        pdf.ln(4)

    def h2(text):
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(51, 65, 85)
        pdf.multi_cell(0, 7, clean(text))
        pdf.ln(1)

    def body(text):
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(55, 65, 81)
        pdf.multi_cell(0, 6, clean(text))
        pdf.ln(2)

    def bullet(text):
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(55, 65, 81)
        pdf.set_x(pdf.l_margin + 4)
        pdf.multi_cell(0, 6, f"-  {clean(text)}")

    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(27, 46, 75)
    pdf.multi_cell(0, 10, clean(prd.title))
    pdf.ln(2)
    pdf.set_font("Helvetica", "I", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.multi_cell(0, 6, clean(prd.tldr))
    pdf.ln(6)

    h1("Problem Statement")
    body(prd.problem_statement)
    h1("Goal")
    body(prd.goal)
    h1("Target Users")
    for u in prd.target_users:
        bullet(u)
    pdf.ln(2)

    if prd.competitive_context:
        h1("Competitive Context")
        for comp in prd.competitive_context:
            h2(comp.competitor)
            body("How they solve it: " + comp.how_they_solve_it)
            body("Gap / opportunity: " + comp.gap_or_opportunity)

    h1("User Stories")
    priority_map = {"must-have": "Must-have", "should-have": "Should-have", "nice-to-have": "Nice-to-have"}
    for i, story in enumerate(prd.user_stories, 1):
        h2(f"Story {i} - {priority_map.get(story.priority, story.priority)}")
        body(f"As a {story.persona}, I want {story.goal} so that {story.benefit}")
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(37, 99, 235)
        pdf.cell(0, 5, "Acceptance Criteria")
        pdf.ln(5)
        for ac in story.criteria:
            bullet(f"Given {ac.given} - when {ac.when} - then {ac.then}")
        pdf.ln(2)

    h1("Success Metrics")
    col_widths = [70, 35, 35, 30]
    headers = ["Metric", "Baseline", "Target", "Timeframe"]
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(241, 245, 249)
    pdf.set_text_color(51, 65, 85)
    for w, h in zip(col_widths, headers):
        pdf.cell(w, 7, h, border=1, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(55, 65, 81)
    for m in prd.success_metrics:
        for w, txt in zip(col_widths, [m.metric, m.baseline, m.target, m.timeframe]):
            pdf.cell(w, 6, clean(str(txt))[:40], border=1)
        pdf.ln()
    pdf.ln(4)

    h1("Out of Scope")
    for o in prd.out_of_scope:
        bullet(o)
    h1("Edge Cases")
    for e in prd.edge_cases:
        bullet(e)
    h1("Risks")
    for r in prd.risks:
        bullet(r)
    if prd.open_questions:
        h1("Open Questions")
        for q in prd.open_questions:
            bullet(f"[ ]  {q}")

    return bytes(pdf.output())
